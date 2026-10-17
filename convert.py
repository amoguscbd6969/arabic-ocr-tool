"""Load common document formats as typed text and/or OCR-ready images.

Task 2 policy:
- Real text is extracted directly; sending it through OCR would reduce accuracy.
- Scanned pages and embedded document photos are returned as PIL images for OCR.
- No permanent conversion directory is created.
"""

from __future__ import annotations

import csv
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pymupdf
import xlrd
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".pttx", ".ptt",
    ".xlsx", ".xls", ".csv", ".tsv", ".txt",
}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS


@dataclass
class DocumentPart:
    """One ordered unit from a source document."""

    kind: Literal["text", "image"]
    content: str | Image.Image
    source: str
    index: int

    def summary(self) -> dict:
        if self.kind == "text":
            return {
                "kind": "text",
                "source": self.source,
                "index": self.index,
                "characters": len(str(self.content)),
            }
        image = self.content
        return {
            "kind": "image",
            "source": self.source,
            "index": self.index,
            "width": image.width,
            "height": image.height,
        }


def _clean_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _pil_from_bytes(data: bytes) -> Image.Image:
    with Image.open(io.BytesIO(data)) as image:
        return image.convert("RGB").copy()


def _render_pdf_page(page, dpi: int = 220) -> Image.Image:
    scale = dpi / 72
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def _large_pdf_images(document, page) -> list[Image.Image]:
    """Extract only meaningful page images, ignoring logos and tiny decorations."""
    page_area = max(1.0, page.rect.width * page.rect.height)
    result = []
    seen = set()
    for image_info in page.get_images(full=True):
        xref = image_info[0]
        if xref in seen:
            continue
        seen.add(xref)
        rectangles = page.get_image_rects(xref)
        coverage = max(
            ((rect.width * rect.height) / page_area for rect in rectangles),
            default=0.0,
        )
        extracted = document.extract_image(xref)
        width = int(extracted.get("width", 0))
        height = int(extracted.get("height", 0))
        if coverage < 0.08 or width < 200 or height < 120:
            continue
        result.append(_pil_from_bytes(extracted["image"]))
    return result


def _load_pdf(path: Path) -> list[DocumentPart]:
    parts = []
    with pymupdf.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = _clean_text(page.get_text("text"))
            page_area = max(1.0, page.rect.width * page.rect.height)
            maximum_coverage = 0.0
            for image_info in page.get_images(full=True):
                for rect in page.get_image_rects(image_info[0]):
                    maximum_coverage = max(
                        maximum_coverage,
                        (rect.width * rect.height) / page_area,
                    )

            source = f"page {page_number}"
            # A full-page raster with a hidden OCR layer is still a scan. OCR the
            # actual pixels instead of trusting potentially bad embedded text.
            if maximum_coverage >= 0.60 or not text:
                parts.append(DocumentPart("image", _render_pdf_page(page), source, page_number))
                continue

            parts.append(DocumentPart("text", text, source, page_number))
            for image_number, image in enumerate(_large_pdf_images(document, page), start=1):
                parts.append(DocumentPart(
                    "image", image, f"{source}, image {image_number}", page_number
                ))
    return parts


def _iter_docx_text(document: Document) -> list[str]:
    lines = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    lines.append("\t".join(values))
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    lines.append(paragraph.text.strip())
    return lines


def _load_docx(path: Path) -> list[DocumentPart]:
    document = Document(path)
    parts = []
    text = _clean_text("\n".join(_iter_docx_text(document)))
    if text:
        parts.append(DocumentPart("text", text, "document text", 1))

    image_number = 0
    seen_parts = set()
    for relationship in document.part.rels.values():
        target = getattr(relationship, "target_part", None)
        content_type = getattr(target, "content_type", "")
        if not content_type.startswith("image/") or target.partname in seen_parts:
            continue
        seen_parts.add(target.partname)
        try:
            image = _pil_from_bytes(target.blob)
        except Exception:
            continue
        if image.width < 200 or image.height < 120:
            continue
        image_number += 1
        parts.append(DocumentPart("image", image, f"embedded image {image_number}", image_number))
    return parts


def _load_pptx(path: Path) -> list[DocumentPart]:
    presentation = Presentation(path)
    parts = []
    part_index = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        image_number = 0
        for shape in sorted(slide.shapes, key=lambda item: (item.top, item.left)):
            if getattr(shape, "has_text_frame", False):
                text = _clean_text(shape.text)
                if text:
                    part_index += 1
                    parts.append(DocumentPart("text", text, f"slide {slide_number}", part_index))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = _pil_from_bytes(shape.image.blob)
                except Exception:
                    continue
                if image.width < 200 or image.height < 120:
                    continue
                image_number += 1
                part_index += 1
                parts.append(DocumentPart(
                    "image",
                    image,
                    f"slide {slide_number}, image {image_number}",
                    part_index,
                ))
    return parts


def _load_xlsx(path: Path) -> list[DocumentPart]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    parts = []
    part_index = 0
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            while values and not values[-1]:
                values.pop()
            if values:
                rows.append("\t".join(values))
        text = _clean_text("\n".join(rows))
        if text:
            part_index += 1
            parts.append(DocumentPart("text", text, f"sheet {worksheet.title}", part_index))
        for image_number, embedded in enumerate(getattr(worksheet, "_images", []), start=1):
            try:
                image = _pil_from_bytes(embedded._data())
            except Exception:
                continue
            if image.width < 200 or image.height < 120:
                continue
            part_index += 1
            parts.append(DocumentPart(
                "image",
                image,
                f"sheet {worksheet.title}, image {image_number}",
                part_index,
            ))
    workbook.close()
    return parts


def _load_xls(path: Path) -> list[DocumentPart]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    parts = []
    for index, worksheet in enumerate(workbook.sheets(), start=1):
        rows = []
        for row_index in range(worksheet.nrows):
            values = [str(worksheet.cell_value(row_index, column)) for column in range(worksheet.ncols)]
            while values and not values[-1]:
                values.pop()
            if values:
                rows.append("\t".join(values))
        text = _clean_text("\n".join(rows))
        if text:
            parts.append(DocumentPart("text", text, f"sheet {worksheet.name}", index))
    workbook.release_resources()
    return parts


def _decode_text_file(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeError(f"Cannot decode text file: {path}")


def _load_delimited(path: Path) -> list[DocumentPart]:
    text = _decode_text_file(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        while row and not row[-1]:
            row.pop()
        if row:
            rows.append("\t".join(cell.strip() for cell in row))
    cleaned = _clean_text("\n".join(rows))
    return [DocumentPart("text", cleaned, "table", 1)] if cleaned else []


def _find_soffice() -> str | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable:
        return executable
    common = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    return str(next((path for path in common if path.exists()), "")) or None


def _legacy_to_modern(path: Path) -> tuple[Path, tempfile.TemporaryDirectory]:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError(
            f"{path.suffix} is a legacy binary format. Install LibreOffice so it can be "
            "converted safely, or save it as DOCX/PPTX first."
        )
    target_extension = ".docx" if path.suffix.lower() == ".doc" else ".pptx"
    temporary = tempfile.TemporaryDirectory(prefix="arabic_ocr_convert_")
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", target_extension[1:], "--outdir", temporary.name, str(path)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    converted = Path(temporary.name) / f"{path.stem}{target_extension}"
    if result.returncode != 0 or not converted.exists():
        temporary.cleanup()
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr or result.stdout}")
    return converted, temporary


def load_document(path: str | Path) -> list[DocumentPart]:
    """Load a supported input while preserving typed text and OCR-worthy images."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported input type {extension!r}. Supported: {supported}")
    if extension in IMAGE_EXTENSIONS:
        with Image.open(path) as image:
            return [DocumentPart("image", image.convert("RGB").copy(), "image", 1)]
    if extension == ".pdf":
        return _load_pdf(path)
    if extension == ".docx":
        return _load_docx(path)
    if extension in {".pptx", ".pttx"}:
        return _load_pptx(path)
    if extension == ".xlsx":
        return _load_xlsx(path)
    if extension == ".xls":
        return _load_xls(path)
    if extension in {".csv", ".tsv"}:
        return _load_delimited(path)
    if extension == ".txt":
        text = _clean_text(_decode_text_file(path))
        return [DocumentPart("text", text, "text file", 1)] if text else []
    if extension in {".doc", ".ppt", ".ptt"}:
        converted, temporary = _legacy_to_modern(path)
        try:
            return _load_docx(converted) if converted.suffix == ".docx" else _load_pptx(converted)
        finally:
            temporary.cleanup()
    raise AssertionError(f"Unhandled extension: {extension}")


def classify_image_in_pdf(pdf_path):
    return [part.summary() for part in _load_pdf(Path(pdf_path))]


def classify_image_in_docx(docx_path):
    return [part.summary() for part in _load_docx(Path(docx_path))]


def classify_image_in_pttx(pttx_path):
    return [part.summary() for part in _load_pptx(Path(pttx_path))]


def classify_image_in_xlsx(xlsx_path):
    return [part.summary() for part in _load_xlsx(Path(xlsx_path))]


def pdf_to_images(pdf_path, dpi: int = 220):
    with pymupdf.open(pdf_path) as document:
        return [_render_pdf_page(page, dpi=dpi) for page in document]


def render_document_pages(path: str | Path, dpi: int = 220) -> list[DocumentPart]:
    """Render every document page for full-page OCR.

    PDF is rendered directly. Office documents are first converted to PDF with
    LibreOffice so pagination and visual layout are preserved.
    """
    path = Path(path)
    extension = path.suffix.lower()
    if extension == ".pdf":
        images = pdf_to_images(path, dpi=dpi)
    elif extension in {".docx", ".doc", ".pptx", ".ppt", ".pttx", ".ptt", ".xlsx", ".xls"}:
        with tempfile.TemporaryDirectory(prefix="arabic_ocr_full_pages_") as temporary:
            pdf_path = Path(temporary) / f"{path.stem}.pdf"
            _office_to_pdf(path, pdf_path)
            images = pdf_to_images(pdf_path, dpi=dpi)
    else:
        raise ValueError(
            f"Full-page OCR is not meaningful for {extension or 'this input type'}; "
            "use mixed mode for text/CSV files."
        )
    return [
        DocumentPart("image", image, f"page {page_number}", page_number)
        for page_number, image in enumerate(images, start=1)
    ]


def images_to_pdf(images, output_pdf_path):
    converted = [image.convert("RGB") for image in images]
    if not converted:
        raise ValueError("images must not be empty")
    converted[0].save(output_pdf_path, "PDF", save_all=True, append_images=converted[1:])


def _office_to_pdf(input_path, output_pdf_path):
    soffice = _find_soffice()
    if not soffice:
        return _microsoft_office_to_pdf(input_path, output_pdf_path)
    input_path, output_pdf_path = Path(input_path), Path(output_pdf_path)
    with tempfile.TemporaryDirectory(prefix="arabic_ocr_pdf_") as temporary:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", temporary, str(input_path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        generated = Path(temporary) / f"{input_path.stem}.pdf"
        if result.returncode != 0 or not generated.exists():
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr or result.stdout}")
        output_pdf_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(generated, output_pdf_path)


def _microsoft_office_to_pdf(input_path, output_pdf_path):
    """Render an Office document with an installed Microsoft Office app."""
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "Full OCR for Office files requires LibreOffice, or Microsoft Office "
            "together with pywin32."
        ) from exc

    input_path = Path(input_path).resolve()
    output_pdf_path = Path(output_pdf_path).resolve()
    output_pdf_path.parent.mkdir(parents=True, exist_ok=True)
    extension = input_path.suffix.lower()
    pythoncom.CoInitialize()
    application = None
    document = None
    try:
        if extension in {".doc", ".docx"}:
            application = win32com.client.DispatchEx("Word.Application")
            application.Visible = False
            application.DisplayAlerts = 0
            document = application.Documents.Open(str(input_path), ReadOnly=True)
            document.ExportAsFixedFormat(str(output_pdf_path), 17)
        elif extension in {".ppt", ".pptx", ".ptt", ".pttx"}:
            application = win32com.client.DispatchEx("PowerPoint.Application")
            document = application.Presentations.Open(
                str(input_path), ReadOnly=True, Untitled=False, WithWindow=False
            )
            document.SaveAs(str(output_pdf_path), 32)
        elif extension in {".xls", ".xlsx"}:
            application = win32com.client.DispatchEx("Excel.Application")
            application.Visible = False
            application.DisplayAlerts = False
            document = application.Workbooks.Open(str(input_path), ReadOnly=True)
            document.ExportAsFixedFormat(0, str(output_pdf_path))
        else:
            raise ValueError(f"Unsupported Office format for page rendering: {extension}")
    except Exception as exc:
        raise RuntimeError(f"Microsoft Office PDF rendering failed for {input_path}: {exc}") from exc
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

    if not output_pdf_path.is_file():
        raise RuntimeError(f"Microsoft Office did not create {output_pdf_path}")


def docx_to_pdf(docx_path, output_pdf_path):
    _office_to_pdf(docx_path, output_pdf_path)


def pttx_to_pdf(pttx_path, output_pdf_path):
    _office_to_pdf(pttx_path, output_pdf_path)


def pptx_to_pdf(pptx_path, output_pdf_path):
    _office_to_pdf(pptx_path, output_pdf_path)


def xlsx_to_pdf(xlsx_path, output_pdf_path):
    _office_to_pdf(xlsx_path, output_pdf_path)


def pdf_to_docx(pdf_path, output_docx_path):
    document = Document()
    for index, image in enumerate(pdf_to_images(pdf_path), start=1):
        if index > 1:
            document.add_page_break()
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        document.add_picture(stream, width=Inches(6.3))
    document.save(output_docx_path)


def pdf_to_pttx(pdf_path, output_pttx_path):
    presentation = Presentation()
    presentation.slide_width = Inches(8.27)
    presentation.slide_height = Inches(11.69)
    blank_layout = presentation.slide_layouts[6]
    for image in pdf_to_images(pdf_path):
        slide = presentation.slides.add_slide(blank_layout)
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        slide.shapes.add_picture(
            stream, 0, 0, width=presentation.slide_width, height=presentation.slide_height
        )
    presentation.save(output_pttx_path)


def pdf_to_pptx(pdf_path, output_pptx_path):
    pdf_to_pttx(pdf_path, output_pptx_path)


def pdf_to_xlsx(pdf_path, output_xlsx_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    with pymupdf.open(pdf_path) as document:
        for page_number, page in enumerate(document, start=1):
            worksheet = workbook.create_sheet(f"Page {page_number}")
            text = _clean_text(page.get_text("text"))
            if text:
                for row, line in enumerate(text.splitlines(), start=1):
                    worksheet.cell(row=row, column=1, value=line)
            else:
                image = _render_pdf_page(page)
                stream = io.BytesIO()
                image.save(stream, format="PNG")
                stream.seek(0)
                worksheet.add_image(ExcelImage(stream), "A1")
    workbook.save(output_xlsx_path)
